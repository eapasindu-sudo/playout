"""Ad report (proof of play) from the as-run logs: which commercial went on air, when, for which client.

Exports to Excel (.xlsx), PDF and CSV for a date range, optionally for one client.
"""
import csv
import datetime as dt
import io
import os

AIRED = ("played", "skipped", "stopped", "taken off")


def _days(d1, d2):
    d = d1
    while d <= d2:
        yield d
        d += dt.timedelta(days=1)


def collect(asrun, d1, d2, client="", text="", ads_only=True):
    """Rows of the as-run logs between two dates (inclusive), filtered."""
    client = (client or "").strip().lower()
    text = (text or "").strip().lower()
    rows = []
    for day in _days(d1, d2):
        for r in asrun.read(day):
            if ads_only and r.get("type") != "clip":
                continue
            if client and (r.get("client") or "").strip().lower() != client:
                continue
            if text and text not in (r.get("item") or "").lower() and text not in (r.get("file") or "").lower():
                continue
            rows.append(r)
    return rows


def clients(asrun, d1, d2):
    seen = set()
    for day in _days(d1, d2):
        for r in asrun.read(day):
            c = (r.get("client") or "").strip()
            if c:
                seen.add(c)
    return sorted(seen, key=str.lower)


def summarize(rows):
    """Per client + item: plays, total seconds on air, first and last time."""
    out = {}
    for r in rows:
        if not (r.get("status") or "").startswith(AIRED):
            continue
        key = ((r.get("client") or "").strip(), r.get("item") or "")
        s = out.setdefault(key, {"client": key[0], "item": key[1], "plays": 0, "seconds": 0.0,
                                 "first": None, "last": None})
        s["plays"] += 1
        try:
            s["seconds"] += float(r.get("played_sec") or 0)
        except ValueError:
            pass
        when = f"{r.get('date')} {r.get('start')}"
        s["first"] = min(s["first"] or when, when)
        s["last"] = max(s["last"] or when, when)
    return sorted(out.values(), key=lambda s: (s["client"].lower(), s["item"].lower()))


def _hms(sec):
    sec = int(round(sec))
    return f"{sec // 3600:02d}:{sec // 60 % 60:02d}:{sec % 60:02d}"


def _title(station, d1, d2, client):
    period = f"{d1:%d %b %Y}" if d1 == d2 else f"{d1:%d %b %Y} - {d2:%d %b %Y}"
    return (station or "Endora Playout"), f"Commercial air report{' - ' + client if client else ''}", period


DETAIL_COLS = ["Date", "Start", "End", "Client", "Commercial", "Planned", "Aired", "Status"]


def _detail(r):
    return [r.get("date", ""), r.get("start", ""), r.get("end", ""), r.get("client", ""), r.get("item", ""),
            r.get("planned_sec", ""), r.get("played_sec", ""), r.get("status", "")]


def to_csv(rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(DETAIL_COLS)
    for r in rows:
        w.writerow(_detail(r))
    return buf.getvalue().encode("utf-8-sig")      # BOM so Excel opens it correctly


def to_xlsx(rows, station, d1, d2, client):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    st, title, period = _title(station, d1, d2, client)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2B2E32")

    ws = wb.active
    ws.title = "Summary"
    ws.append([st])
    ws.append([title])
    ws.append([period])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(bold=True)
    ws.append([])
    ws.append(["Client", "Commercial", "Plays", "Total air time", "First aired", "Last aired"])
    for c in ws[5]:
        c.font, c.fill = head, fill
    summ = summarize(rows)
    for s in summ:
        ws.append([s["client"], s["item"], s["plays"], _hms(s["seconds"]), s["first"], s["last"]])
    ws.append([])
    ws.append(["Total", "", sum(s["plays"] for s in summ), _hms(sum(s["seconds"] for s in summ))])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for col, w in zip("ABCDEF", (24, 36, 8, 15, 20, 20)):
        ws.column_dimensions[col].width = w

    wd = wb.create_sheet("Spots")
    wd.append(DETAIL_COLS)
    for c in wd[1]:
        c.font, c.fill = head, fill
    for r in rows:
        wd.append(_detail(r))
    for col, w in zip("ABCDEFGH", (12, 10, 10, 22, 36, 9, 9, 14)):
        wd.column_dimensions[col].width = w
    for row in wd.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rows, station, d1, d2, client):
    from fpdf import FPDF

    st, title, period = _title(station, d1, d2, client)

    def clean(s):
        return str(s or "").encode("latin-1", "replace").decode("latin-1")

    class PDF(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", size=8)
            self.set_text_color(120)
            self.cell(0, 6, f"Printed {dt.datetime.now():%d %b %Y %H:%M}  -  page {self.page_no()}/{{nb}}", align="C")

    pdf = PDF(orientation="P", unit="mm", format="A4")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 8, clean(st), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7, clean(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 6, clean(period), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    def table(cols, widths, data, bold_last=False):
        pdf.set_font("Helvetica", "B", 8.5)
        pdf.set_fill_color(43, 46, 50)
        pdf.set_text_color(255)
        for c, w in zip(cols, widths):
            pdf.cell(w, 6, clean(c), border=0, fill=True)
        pdf.ln()
        pdf.set_text_color(0)
        for i, row in enumerate(data):
            last = bold_last and i == len(data) - 1
            pdf.set_font("Helvetica", "B" if last else "", 8.5)
            pdf.set_fill_color(242, 242, 242)
            for v, w in zip(row, widths):
                txt = clean(v)
                while txt and pdf.get_string_width(txt) > w - 1.5:
                    txt = txt[:-1]
                pdf.cell(w, 5.5, txt, fill=i % 2 == 1)
            pdf.ln()

    summ = summarize(rows)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 7, "Summary", new_x="LMARGIN", new_y="NEXT")
    data = [[s["client"], s["item"], s["plays"], _hms(s["seconds"]), s["first"][5:16], s["last"][5:16]] for s in summ]
    data.append(["Total", "", sum(s["plays"] for s in summ), _hms(sum(s["seconds"] for s in summ)), "", ""])
    table(["Client", "Commercial", "Plays", "Air time", "First", "Last"], [36, 64, 14, 20, 23, 23], data, True)
    pdf.ln(5)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 7, "Spot list (as run)", new_x="LMARGIN", new_y="NEXT")
    table(["Date", "Start", "End", "Client", "Commercial", "Aired", "Status"], [20, 15, 15, 32, 64, 14, 20],
          [[r.get("date"), r.get("start"), r.get("end"), r.get("client"), r.get("item"),
            (r.get("played_sec") or "") + "s", r.get("status")] for r in rows])
    return bytes(pdf.output())


def filename(kind, d1, d2, client):
    who = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (client or "all"))
    span = f"{d1:%Y-%m-%d}" if d1 == d2 else f"{d1:%Y-%m-%d}_to_{d2:%Y-%m-%d}"
    return f"AdReport_{who}_{span}.{kind}"
